# -*- coding: utf-8 -*-
"""Generazione dei report mensili in pipeline/output/AAAA-MM/.

Produce:
  - un file Excel per AREA (foglio di sintesi + un foglio per ogni CdC
    dell'area) con actual mese/YTD vs budget e scostamenti;
  - scostamenti.csv con gli scostamenti per CdC (fonte machine-readable
    da citare nei commenti, umani o AI, agli scostamenti);
  - una sintesi HTML per la direzione (totali per area, top 10
    scostamenti sfavorevoli);
  - riepilogo_quadrature.txt con i totali di controllo pre/post
    ribaltamento.

Convenzione di segno: i valori sono COSTI, quindi uno scostamento
positivo (actual > budget) e' SFAVOREVOLE e viene evidenziato in rosso.
Gli scostamenti sono calcolati sui costi diretti (il budget e' per
CdC/natura sui costi diretti); i ribaltamenti sono esposti a parte
come "costo pieno".
"""
from datetime import datetime

import pandas as pd
from jinja2 import Template
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ingest import OUTPUT_DIR

CLASSI_RIBALTABILI = ("ausiliario", "struttura")

# ----------------------------------------------------------- stili openpyxl
FILL_INTESTAZIONE = PatternFill("solid", fgColor="1F4E79")
FONT_INTESTAZIONE = Font(bold=True, color="FFFFFF")
FILL_TOTALE = PatternFill("solid", fgColor="D9E2F3")
FONT_TOTALE = Font(bold=True)
FONT_ROSSO = Font(color="C00000", bold=True)
FONT_TITOLO = Font(bold=True, size=14)
FONT_NOTA = Font(italic=True, color="808080")
FMT_EUR = "#,##0.00"
FMT_PCT = "0.0%"


def eur(cent):
    """Formatta centesimi come importo in stile italiano: 1.234.567,89."""
    testo = f"{cent / 100:,.2f}"
    return testo.replace(",", "§").replace(".", ",").replace("§", ".")


def pct(delta_cent, budget_cent):
    """Scostamento percentuale sul budget, formato italiano."""
    if budget_cent == 0:
        return "n.d."
    return f"{delta_cent / budget_cent * 100:.1f}%".replace(".", ",")


def _quota_pct(delta_cent, budget_cent):
    """Scostamento percentuale come frazione (per i formati Excel)."""
    return delta_cent / budget_cent if budget_cent else None


# ================================================================ dati
def _somma_natura(df, nome):
    g = df.groupby(["cdc", "natura_costo"], as_index=False)["importo_cent"].sum()
    return g.rename(columns={"importo_cent": nome, "cdc": "codice"})


def _somma_cdc(df, mese, nome_mese, nome_ytd):
    m = (df[df["periodo"] == mese].groupby("cdc", as_index=False)
         ["importo_cent"].sum().rename(columns={"importo_cent": nome_mese}))
    y = (df.groupby("cdc", as_index=False)["importo_cent"].sum()
         .rename(columns={"importo_cent": nome_ytd}))
    return m.merge(y, on="cdc", how="outer").rename(columns={"cdc": "codice"})


def _prepara_dati(mese, config, esito, budget):
    """Costruisce le tabelle actual/budget per natura, CdC e area."""
    cdc = config["cdc"]
    foglie = cdc[(cdc["livello"] == 3) & (cdc["attivo"] == "SI")].copy()
    diretti = esito["diretti"]
    post = esito["post"]

    # dettaglio per CdC / natura (costi diretti vs budget)
    nat = _somma_natura(diretti[diretti["periodo"] == mese], "act_mese")
    for frame, nome in ((diretti, "act_ytd"),
                        (budget[budget["periodo"] == mese], "bud_mese"),
                        (budget, "bud_ytd")):
        nat = nat.merge(_somma_natura(frame, nome),
                        on=["codice", "natura_costo"], how="outer")
    nat = nat.fillna(0)
    for col in ("act_mese", "act_ytd", "bud_mese", "bud_ytd"):
        nat[col] = nat[col].astype("int64")

    # totali per CdC: diretti, ribaltamenti ricevuti, costo pieno
    tot = nat.groupby("codice", as_index=False)[
        ["act_mese", "act_ytd", "bud_mese", "bud_ytd"]].sum()
    ricevuti = _somma_cdc(post[post["origine_ribaltamento"] != "DIRETTO"],
                          mese, "ric_mese", "ric_ytd")
    pieno = _somma_cdc(post, mese, "pieno_mese", "pieno_ytd")

    per_cdc = (foglie[["codice", "descrizione", "classe", "area",
                       "responsabile"]]
               .merge(tot, on="codice", how="left")
               .merge(ricevuti, on="codice", how="left")
               .merge(pieno, on="codice", how="left")
               .fillna(0))
    for col in ("act_mese", "act_ytd", "bud_mese", "bud_ytd", "ric_mese",
                "ric_ytd", "pieno_mese", "pieno_ytd"):
        per_cdc[col] = per_cdc[col].astype("int64")

    # per i CdC ribaltati i costi diretti vengono interamente ceduti
    ribaltato = per_cdc["classe"].isin(CLASSI_RIBALTABILI)
    per_cdc["ced_mese"] = per_cdc["act_mese"].where(ribaltato, 0)
    per_cdc["ced_ytd"] = per_cdc["act_ytd"].where(ribaltato, 0)
    per_cdc["ribalt_mese"] = per_cdc["ric_mese"] - per_cdc["ced_mese"]
    per_cdc["delta_mese"] = per_cdc["act_mese"] - per_cdc["bud_mese"]
    per_cdc["delta_ytd"] = per_cdc["act_ytd"] - per_cdc["bud_ytd"]

    colonne_num = ["act_mese", "bud_mese", "delta_mese", "act_ytd", "bud_ytd",
                   "delta_ytd", "ribalt_mese", "pieno_mese", "pieno_ytd"]
    per_area = per_cdc.groupby("area", as_index=False)[colonne_num].sum()

    top10 = (per_cdc[per_cdc["delta_mese"] > 0]
             .sort_values("delta_mese", ascending=False).head(10))

    nat_per_cdc = {codice: gruppo.sort_values("bud_ytd", ascending=False)
                   for codice, gruppo in nat.groupby("codice")}

    return {"foglie": foglie, "nat_per_cdc": nat_per_cdc,
            "per_cdc": per_cdc, "per_area": per_area, "top10": top10}


# ================================================================ Excel
def _intestazione(ws, riga, titoli):
    for i, titolo in enumerate(titoli, start=1):
        cella = ws.cell(row=riga, column=i, value=titolo)
        cella.fill = FILL_INTESTAZIONE
        cella.font = FONT_INTESTAZIONE
        cella.alignment = Alignment(horizontal="center", wrap_text=True)


def _numero(ws, riga, col, cent, rosso=False, bold=False):
    cella = ws.cell(row=riga, column=col, value=round(cent / 100, 2))
    cella.number_format = FMT_EUR
    if rosso:
        cella.font = FONT_ROSSO
    elif bold:
        cella.font = FONT_TOTALE
    return cella


def _percento(ws, riga, col, delta_cent, budget_cent, rosso=False):
    valore = _quota_pct(delta_cent, budget_cent)
    cella = ws.cell(row=riga, column=col, value=valore)
    cella.number_format = FMT_PCT
    if rosso and valore is not None:
        cella.font = FONT_ROSSO
    return cella


def _foglio_sintesi_area(ws, area, df_area, mese):
    ws.cell(row=1, column=1, value=f"Sintesi area {area} — {mese}") \
        .font = FONT_TITOLO
    ws.cell(row=2, column=1,
            value="DATI SINTETICI DIMOSTRATIVI — importi in EUR; scostamenti "
                  "su costi diretti (actual - budget), positivi = sfavorevoli") \
        .font = FONT_NOTA
    _intestazione(ws, 4, ["CdC", "Descrizione", "Classe", "Actual mese",
                          "Budget mese", "Scost. mese", "Scost. %",
                          "Actual YTD", "Budget YTD", "Scost. YTD",
                          "Ribaltamenti mese (+ric/-ced)",
                          "Costo pieno mese"])
    riga = 5
    for r in df_area.sort_values("codice").itertuples(index=False):
        ws.cell(row=riga, column=1, value=r.codice)
        ws.cell(row=riga, column=2, value=r.descrizione)
        ws.cell(row=riga, column=3, value=r.classe)
        sfav_m, sfav_y = r.delta_mese > 0, r.delta_ytd > 0
        _numero(ws, riga, 4, r.act_mese)
        _numero(ws, riga, 5, r.bud_mese)
        _numero(ws, riga, 6, r.delta_mese, rosso=sfav_m)
        _percento(ws, riga, 7, r.delta_mese, r.bud_mese, rosso=sfav_m)
        _numero(ws, riga, 8, r.act_ytd)
        _numero(ws, riga, 9, r.bud_ytd)
        _numero(ws, riga, 10, r.delta_ytd, rosso=sfav_y)
        _numero(ws, riga, 11, r.ribalt_mese)
        _numero(ws, riga, 12, r.pieno_mese)
        riga += 1
    # riga di totale area
    ws.cell(row=riga, column=1, value="TOTALE AREA").font = FONT_TOTALE
    somme = df_area[["act_mese", "bud_mese", "delta_mese", "act_ytd",
                     "bud_ytd", "delta_ytd", "ribalt_mese",
                     "pieno_mese"]].sum()
    colonne = {4: "act_mese", 5: "bud_mese", 6: "delta_mese", 8: "act_ytd",
               9: "bud_ytd", 10: "delta_ytd", 11: "ribalt_mese",
               12: "pieno_mese"}
    for col, nome in colonne.items():
        cella = _numero(ws, riga, col, int(somme[nome]), bold=True)
        cella.fill = FILL_TOTALE
    _percento(ws, riga, 7, int(somme["delta_mese"]), int(somme["bud_mese"]))
    for col in (1, 2, 3, 7):
        ws.cell(row=riga, column=col).fill = FILL_TOTALE
    # larghezze colonne e blocco intestazioni
    larghezze = [14, 34, 12, 14, 14, 13, 9, 14, 14, 13, 16, 14]
    for i, larghezza in enumerate(larghezze, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larghezza
    ws.freeze_panes = "A5"


def _foglio_cdc(ws, info, dettaglio, mese):
    """Foglio di dettaglio di un singolo CdC: nature vs budget + ribaltamenti."""
    ws.cell(row=1, column=1,
            value=f"{info.codice} — {info.descrizione}").font = FONT_TITOLO
    ws.cell(row=2, column=1,
            value=f"Area: {info.area}  |  Classe: {info.classe}  |  "
                  f"Responsabile: {info.responsabile}").font = FONT_NOTA
    ws.cell(row=3, column=1,
            value=f"Periodo: {mese}  |  Importi in EUR  |  "
                  f"Dati sintetici dimostrativi").font = FONT_NOTA
    _intestazione(ws, 5, ["Natura di costo", "Actual mese", "Budget mese",
                          "Scost. mese", "Scost. %", "Actual YTD",
                          "Budget YTD", "Scost. YTD", "Scost. % YTD"])
    riga = 6
    if dettaglio is None:
        dettaglio = pd.DataFrame(columns=["natura_costo", "act_mese",
                                          "bud_mese", "act_ytd", "bud_ytd"])
    for r in dettaglio.itertuples(index=False):
        delta_m = r.act_mese - r.bud_mese
        delta_y = r.act_ytd - r.bud_ytd
        ws.cell(row=riga, column=1, value=r.natura_costo)
        _numero(ws, riga, 2, r.act_mese)
        _numero(ws, riga, 3, r.bud_mese)
        _numero(ws, riga, 4, delta_m, rosso=delta_m > 0)
        _percento(ws, riga, 5, delta_m, r.bud_mese, rosso=delta_m > 0)
        _numero(ws, riga, 6, r.act_ytd)
        _numero(ws, riga, 7, r.bud_ytd)
        _numero(ws, riga, 8, delta_y, rosso=delta_y > 0)
        _percento(ws, riga, 9, delta_y, r.bud_ytd, rosso=delta_y > 0)
        riga += 1
    # totale costi diretti
    ws.cell(row=riga, column=1, value="Totale costi diretti").font = FONT_TOTALE
    delta_m = info.act_mese - info.bud_mese
    delta_y = info.act_ytd - info.bud_ytd
    for col, cent in ((2, info.act_mese), (3, info.bud_mese), (4, delta_m),
                      (6, info.act_ytd), (7, info.bud_ytd), (8, delta_y)):
        cella = _numero(ws, riga, col, cent, bold=True)
        cella.fill = FILL_TOTALE
        if col in (4, 8) and cent > 0:
            cella.font = FONT_ROSSO
    _percento(ws, riga, 5, delta_m, info.bud_mese)
    _percento(ws, riga, 9, delta_y, info.bud_ytd)
    for col in (1, 5, 9):
        ws.cell(row=riga, column=col).fill = FILL_TOTALE
    riga += 1
    # sezione ribaltamenti
    if info.classe in CLASSI_RIBALTABILI:
        ws.cell(row=riga, column=1,
                value="Costi trasferiti con ribaltamento")
        _numero(ws, riga, 2, -info.ced_mese)
        _numero(ws, riga, 6, -info.ced_ytd)
        riga += 1
        ws.cell(row=riga, column=1,
                value="Residuo dopo ribaltamento").font = FONT_TOTALE
        _numero(ws, riga, 2, info.act_mese - info.ced_mese, bold=True)
        _numero(ws, riga, 6, info.act_ytd - info.ced_ytd, bold=True)
    else:
        ws.cell(row=riga, column=1,
                value="Costi ricevuti da ribaltamento")
        _numero(ws, riga, 2, info.ric_mese)
        _numero(ws, riga, 6, info.ric_ytd)
        riga += 1
        ws.cell(row=riga, column=1,
                value="Costo pieno del CdC").font = FONT_TOTALE
        _numero(ws, riga, 2, info.pieno_mese, bold=True)
        _numero(ws, riga, 6, info.pieno_ytd, bold=True)
    larghezze = [30, 13, 13, 12, 9, 13, 13, 12, 10]
    for i, larghezza in enumerate(larghezze, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larghezza
    ws.freeze_panes = "A6"


def _scrivi_excel_area(area, dati, mese, cartella):
    df_area = dati["per_cdc"][dati["per_cdc"]["area"] == area]
    wb = Workbook()
    ws = wb.active
    ws.title = "Sintesi"
    _foglio_sintesi_area(ws, area, df_area, mese)
    for info in df_area.sort_values("codice").itertuples(index=False):
        ws_cdc = wb.create_sheet(title=info.codice[:31])
        _foglio_cdc(ws_cdc, info, dati["nat_per_cdc"].get(info.codice), mese)
    percorso = cartella / f"report_{area.replace(' ', '_')}_{mese}.xlsx"
    wb.save(percorso)
    return percorso


# ================================================================ HTML
TEMPLATE_HTML = Template("""<!doctype html>
<html lang="it">
<head>
<meta charset="utf-8">
<title>Sintesi controllo di gestione {{ mese }}</title>
<style>
  body { font-family: "Segoe UI", Arial, sans-serif; margin: 2rem auto;
         max-width: 70rem; color: #1a1a2e; padding: 0 1rem; }
  .banner { background: #fff3cd; border: 2px solid #b8860b; color: #7a5c00;
            font-weight: bold; text-align: center; padding: 0.8rem;
            border-radius: 6px; margin-bottom: 1.5rem; }
  h1 { color: #1f4e79; } h2 { color: #1f4e79; margin-top: 2rem; }
  .meta { color: #666; font-size: 0.9rem; }
  .kpi { display: flex; gap: 1rem; flex-wrap: wrap; margin: 1rem 0; }
  .kpi div { background: #f0f4fa; border-radius: 8px; padding: 0.8rem 1.2rem; }
  .kpi .v { font-size: 1.3rem; font-weight: bold; display: block; }
  table { border-collapse: collapse; width: 100%; font-size: 0.9rem; }
  th { background: #1f4e79; color: white; padding: 0.4rem 0.6rem;
       text-align: left; }
  td { border-bottom: 1px solid #d8dee9; padding: 0.35rem 0.6rem; }
  td.n, th.n { text-align: right; font-variant-numeric: tabular-nums; }
  tr:nth-child(even) { background: #f7f9fc; }
  .sfav { color: #c00000; font-weight: bold; }
  .fav { color: #1e7a1e; }
  .ok { color: #1e7a1e; font-weight: bold; }
  ul.warn li { color: #7a5c00; }
  footer { margin-top: 2rem; color: #999; font-size: 0.8rem; }
</style>
</head>
<body>
<div class="banner">DATI SINTETICI DIMOSTRATIVI — questo report è generato da
dati fittizi creati ad hoc: nessun valore rappresenta la contabilità reale.</div>
<h1>Controllo di gestione — Sintesi mensile {{ mese }}</h1>
<p class="meta">Generato il {{ generato_il }} dalla pipeline dimostrativa
(actual e scostamenti su costi diretti; ribaltamenti esposti come costo pieno).</p>

<div class="kpi">
  <div><span class="v">{{ kpi.act_mese }} €</span>Costi actual del mese</div>
  <div><span class="v">{{ kpi.bud_mese }} €</span>Budget del mese</div>
  <div><span class="v {{ kpi.classe_delta }}">{{ kpi.delta_mese }} €</span>
       Scostamento mese ({{ kpi.delta_pct }})</div>
  <div><span class="v">{{ kpi.act_ytd }} €</span>Costi actual YTD</div>
  <div><span class="v">{{ kpi.n_cdc }}</span>CdC attivi</div>
  <div><span class="v">{{ kpi.n_movimenti }}</span>Movimenti nel mese</div>
</div>

<h2>Riepilogo per area</h2>
<table>
<tr><th>Area</th><th class="n">Actual mese</th><th class="n">Budget mese</th>
<th class="n">Scost. mese</th><th class="n">Scost. %</th>
<th class="n">Actual YTD</th><th class="n">Budget YTD</th>
<th class="n">Scost. YTD</th><th class="n">Costo pieno mese</th></tr>
{% for a in aree %}
<tr><td>{{ a.area }}</td><td class="n">{{ a.act_mese }}</td>
<td class="n">{{ a.bud_mese }}</td>
<td class="n {{ a.classe_delta }}">{{ a.delta_mese }}</td>
<td class="n {{ a.classe_delta }}">{{ a.delta_pct }}</td>
<td class="n">{{ a.act_ytd }}</td><td class="n">{{ a.bud_ytd }}</td>
<td class="n {{ a.classe_delta_ytd }}">{{ a.delta_ytd }}</td>
<td class="n">{{ a.pieno_mese }}</td></tr>
{% endfor %}
<tr><td><b>TOTALE</b></td><td class="n"><b>{{ tot.act_mese }}</b></td>
<td class="n"><b>{{ tot.bud_mese }}</b></td>
<td class="n {{ tot.classe_delta }}"><b>{{ tot.delta_mese }}</b></td>
<td class="n {{ tot.classe_delta }}"><b>{{ tot.delta_pct }}</b></td>
<td class="n"><b>{{ tot.act_ytd }}</b></td>
<td class="n"><b>{{ tot.bud_ytd }}</b></td>
<td class="n {{ tot.classe_delta_ytd }}"><b>{{ tot.delta_ytd }}</b></td>
<td class="n"><b>{{ tot.pieno_mese }}</b></td></tr>
</table>

<h2>Top 10 scostamenti sfavorevoli del mese (costi diretti vs budget)</h2>
<table>
<tr><th>#</th><th>CdC</th><th>Descrizione</th><th>Area</th>
<th class="n">Actual mese</th><th class="n">Budget mese</th>
<th class="n">Scostamento</th><th class="n">Scost. %</th></tr>
{% for r in top10 %}
<tr><td>{{ loop.index }}</td><td>{{ r.codice }}</td>
<td>{{ r.descrizione }}</td><td>{{ r.area }}</td>
<td class="n">{{ r.act_mese }}</td><td class="n">{{ r.bud_mese }}</td>
<td class="n sfav">{{ r.delta_mese }}</td>
<td class="n sfav">{{ r.delta_pct }}</td></tr>
{% endfor %}
</table>

<h2>Quadrature</h2>
<table>
<tr><th>Controllo</th><th class="n">Valore</th><th class="n">Riscontro</th>
<th>Esito</th></tr>
<tr><td>Totale movimenti del mese vs totale di controllo ERP</td>
<td class="n">{{ q.mov_mese }}</td><td class="n">{{ q.ctrl_mese }}</td>
<td class="ok">{{ q.esito_ctrl }}</td></tr>
<tr><td>Totale costi pre-ribaltamento vs post-ribaltamento (YTD)</td>
<td class="n">{{ q.pre }}</td><td class="n">{{ q.post }}</td>
<td class="ok">{{ q.esito_ribalt }}</td></tr>
</table>

{% if warnings %}
<h2>Segnalazioni della validazione (non bloccanti)</h2>
<ul class="warn">
{% for w in warnings %}<li><b>{{ w.controllo }}</b>: {{ w.esito }}</li>{% endfor %}
</ul>
{% endif %}

<footer>Progetto-AI-CdC — pipeline dimostrativa di controllo di gestione.
Report Excel per area disponibili nella stessa cartella.</footer>
</body>
</html>
""")


def _riga_area_html(r):
    return {"area": r.area,
            "act_mese": eur(r.act_mese), "bud_mese": eur(r.bud_mese),
            "delta_mese": eur(r.delta_mese),
            "delta_pct": pct(r.delta_mese, r.bud_mese),
            "act_ytd": eur(r.act_ytd), "bud_ytd": eur(r.bud_ytd),
            "delta_ytd": eur(r.delta_ytd), "pieno_mese": eur(r.pieno_mese),
            "classe_delta": "sfav" if r.delta_mese > 0 else "fav",
            "classe_delta_ytd": "sfav" if r.delta_ytd > 0 else "fav"}


def _scrivi_html(mese, dati, esito_ribalta, rapporto, totali_controllo,
                 cartella):
    per_cdc = dati["per_cdc"]
    somme = per_cdc[["act_mese", "bud_mese", "act_ytd", "bud_ytd",
                     "pieno_mese"]].sum()
    delta_mese = int(somme["act_mese"] - somme["bud_mese"])
    delta_ytd = int(somme["act_ytd"] - somme["bud_ytd"])

    kpi = {"act_mese": eur(int(somme["act_mese"])),
           "bud_mese": eur(int(somme["bud_mese"])),
           "delta_mese": eur(delta_mese),
           "delta_pct": pct(delta_mese, int(somme["bud_mese"])),
           "classe_delta": "sfav" if delta_mese > 0 else "fav",
           "act_ytd": eur(int(somme["act_ytd"])),
           "n_cdc": len(per_cdc),
           "n_movimenti": rapporto["n_movimenti_mese"]}

    tot = {"act_mese": eur(int(somme["act_mese"])),
           "bud_mese": eur(int(somme["bud_mese"])),
           "delta_mese": eur(delta_mese),
           "delta_pct": pct(delta_mese, int(somme["bud_mese"])),
           "act_ytd": eur(int(somme["act_ytd"])),
           "bud_ytd": eur(int(somme["bud_ytd"])),
           "delta_ytd": eur(delta_ytd),
           "pieno_mese": eur(int(somme["pieno_mese"])),
           "classe_delta": "sfav" if delta_mese > 0 else "fav",
           "classe_delta_ytd": "sfav" if delta_ytd > 0 else "fav"}

    top10 = [{"codice": r.codice, "descrizione": r.descrizione,
              "area": r.area, "act_mese": eur(r.act_mese),
              "bud_mese": eur(r.bud_mese), "delta_mese": eur(r.delta_mese),
              "delta_pct": pct(r.delta_mese, r.bud_mese)}
             for r in dati["top10"].itertuples(index=False)]

    riga_ctrl = totali_controllo[totali_controllo["periodo"] == mese]
    ctrl_cent = int(riga_ctrl["totale_cent"].iloc[0]) if not riga_ctrl.empty else 0
    q = {"mov_mese": eur(rapporto["totale_mese_cent"]) + " €",
         "ctrl_mese": eur(ctrl_cent) + " €",
         "esito_ctrl": "QUADRA" if abs(rapporto["totale_mese_cent"]
                                       - ctrl_cent) <= 1 else "NON QUADRA",
         "pre": eur(esito_ribalta["tot_pre_cent"]) + " €",
         "post": eur(esito_ribalta["tot_post_cent"]) + " €",
         "esito_ribalt": "QUADRA" if esito_ribalta["tot_pre_cent"]
                         == esito_ribalta["tot_post_cent"] else "NON QUADRA"}

    html = TEMPLATE_HTML.render(
        mese=mese,
        generato_il=datetime.now().strftime("%d/%m/%Y %H:%M"),
        kpi=kpi, tot=tot,
        aree=[_riga_area_html(r)
              for r in dati["per_area"].itertuples(index=False)],
        top10=top10, q=q, warnings=rapporto["warning"])

    percorso = cartella / f"sintesi_{mese}.html"
    percorso.write_text(html, encoding="utf-8")
    return percorso


# ============================================================ scostamenti CSV
def _scrivi_scostamenti(mese, dati, cartella):
    """Esporta gli scostamenti per CdC in scostamenti.csv.

    File machine-readable (separatore ',', decimali '.'): e' la fonte da
    citare nei commenti agli scostamenti (es. "[scostamenti.csv, CdC
    S1-PRD-001]"), ordinata dal piu' sfavorevole al piu' favorevole.
    """
    per_cdc = dati["per_cdc"].sort_values("delta_mese", ascending=False)

    def in_euro(colonna):
        return (per_cdc[colonna] / 100).round(2)

    def in_pct(col_delta, col_budget):
        return [round(d / b * 100, 1) if b else None
                for d, b in zip(per_cdc[col_delta], per_cdc[col_budget])]

    df = pd.DataFrame({
        "periodo": mese,
        "cdc": per_cdc["codice"],
        "descrizione": per_cdc["descrizione"],
        "area": per_cdc["area"],
        "classe": per_cdc["classe"],
        "responsabile": per_cdc["responsabile"],
        "actual_mese_eur": in_euro("act_mese"),
        "budget_mese_eur": in_euro("bud_mese"),
        "scostamento_mese_eur": in_euro("delta_mese"),
        "scostamento_mese_pct": in_pct("delta_mese", "bud_mese"),
        "actual_ytd_eur": in_euro("act_ytd"),
        "budget_ytd_eur": in_euro("bud_ytd"),
        "scostamento_ytd_eur": in_euro("delta_ytd"),
        "scostamento_ytd_pct": in_pct("delta_ytd", "bud_ytd"),
        "ribaltamenti_mese_eur": in_euro("ribalt_mese"),
        "costo_pieno_mese_eur": in_euro("pieno_mese"),
    })
    percorso = cartella / "scostamenti.csv"
    df.to_csv(percorso, index=False, float_format="%.2f")
    return percorso


# ================================================================ quadrature
def _scrivi_quadrature(mese, config, movimenti, esito, totali_controllo,
                       rapporto, cartella):
    cdc = config["cdc"]
    classe_cdc = dict(zip(cdc["codice"], cdc["classe"]))
    diretti = esito["diretti"]
    post = esito["post"]

    righe = []
    righe.append(f"RIEPILOGO QUADRATURE — periodo {mese}")
    righe.append(f"Generato il {datetime.now():%d/%m/%Y %H:%M} — "
                 f"DATI SINTETICI DIMOSTRATIVI")
    righe.append("=" * 72)

    righe.append("")
    righe.append("1) Quadratura movimenti caricati vs totale di controllo ERP")
    righe.append(f"   {'Periodo':<10}{'Movimenti':>12}{'Caricato (EUR)':>20}"
                 f"{'Controllo (EUR)':>20}{'Delta':>10}")
    for periodo in sorted(movimenti["periodo"].unique()):
        caricato = int(movimenti.loc[movimenti["periodo"] == periodo,
                                     "importo_cent"].sum())
        n_mov = int((movimenti["periodo"] == periodo).sum())
        riga_ctrl = totali_controllo[totali_controllo["periodo"] == periodo]
        atteso = int(riga_ctrl["totale_cent"].iloc[0]) \
            if not riga_ctrl.empty else 0
        esito_txt = "OK" if abs(caricato - atteso) <= 1 else "KO"
        righe.append(f"   {periodo:<10}{n_mov:>12}{eur(caricato):>20}"
                     f"{eur(atteso):>20}{eur(caricato - atteso):>8} {esito_txt}")

    righe.append("")
    righe.append("2) Quadratura ribaltamento (il totale costi non cambia)")
    pre_mese = int(diretti.loc[diretti["periodo"] == mese,
                               "importo_cent"].sum())
    post_mese = int(post.loc[post["periodo"] == mese, "importo_cent"].sum())
    righe.append(f"   Mese {mese}:")
    righe.append(f"     Totale pre-ribaltamento : {eur(pre_mese):>18} EUR")
    righe.append(f"     Totale post-ribaltamento: {eur(post_mese):>18} EUR")
    righe.append(f"     Delta                   : {eur(post_mese - pre_mese):>18} EUR "
                 f"-> {'OK' if pre_mese == post_mese else 'KO'}")
    righe.append(f"   Cumulato YTD (tutti i periodi caricati):")
    righe.append(f"     Totale pre-ribaltamento : "
                 f"{eur(esito['tot_pre_cent']):>18} EUR")
    righe.append(f"     Totale post-ribaltamento: "
                 f"{eur(esito['tot_post_cent']):>18} EUR")
    righe.append(f"     Delta                   : "
                 f"{eur(esito['tot_post_cent'] - esito['tot_pre_cent']):>18} EUR "
                 f"-> {'OK' if esito['tot_pre_cent'] == esito['tot_post_cent'] else 'KO'}")

    righe.append("")
    righe.append(f"3) Costi del mese {mese} per classe di CdC (pre -> post)")
    pre_classe = (diretti[diretti["periodo"] == mese]
                  .assign(classe=lambda d: d["cdc"].map(classe_cdc))
                  .groupby("classe")["importo_cent"].sum())
    post_classe = (post[post["periodo"] == mese]
                   .assign(classe=lambda d: d["cdc"].map(classe_cdc))
                   .groupby("classe")["importo_cent"].sum())
    for classe in ("produttivo", "ausiliario", "struttura", "commerciale"):
        pre_c = int(pre_classe.get(classe, 0))
        post_c = int(post_classe.get(classe, 0))
        righe.append(f"   {classe:<13}{eur(pre_c):>20} EUR  ->  "
                     f"{eur(post_c):>20} EUR")
    righe.append("   (dopo il ribaltamento i CdC ausiliari e di struttura "
                 "sono a zero)")

    righe.append("")
    righe.append("4) Esito validazione")
    righe.append(f"   Errori bloccanti: {len(rapporto['errori'])}")
    righe.append(f"   Warning        : {len(rapporto['warning'])}")
    for w in rapporto["warning"]:
        righe.append(f"     - {w['controllo']}: {w['esito']}")

    percorso = cartella / "riepilogo_quadrature.txt"
    percorso.write_text("\n".join(righe) + "\n", encoding="utf-8")
    return percorso


# ================================================================ entrata
def genera_report(mese, config, movimenti, esito_ribalta, budget,
                  totali_controllo, rapporto):
    """Genera tutti i report del mese e ritorna l'elenco dei percorsi."""
    cartella = OUTPUT_DIR / mese
    cartella.mkdir(parents=True, exist_ok=True)
    dati = _prepara_dati(mese, config, esito_ribalta, budget)

    percorsi = []
    for area in dati["per_cdc"]["area"].drop_duplicates().sort_values():
        percorsi.append(_scrivi_excel_area(area, dati, mese, cartella))
    percorsi.append(_scrivi_scostamenti(mese, dati, cartella))
    percorsi.append(_scrivi_html(mese, dati, esito_ribalta, rapporto,
                                 totali_controllo, cartella))
    percorsi.append(_scrivi_quadrature(mese, config, movimenti, esito_ribalta,
                                       totali_controllo, rapporto, cartella))
    return percorsi


def main():
    """Entrypoint CLI: rigenera i soli report del mese indicato.

    Riesegue in memoria ingest, validazione (bloccante) e ribaltamento,
    senza rigenerare i dati di esempio.
    """
    import argparse
    import sys
    from pathlib import Path

    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import ingest
    import ribalta
    import valida

    parser = argparse.ArgumentParser(
        description="Rigenera i soli report del mese indicato (riesegue in "
                    "memoria ingest, validazione e ribaltamento).")
    parser.add_argument("--mese", required=True, metavar="YYYY-MM",
                        help="periodo da elaborare, es. 2026-06")
    args = parser.parse_args()
    mese = args.mese

    config = ingest.carica_config()
    movimenti = ingest.carica_movimenti(mese)
    budget = ingest.carica_budget(mese[:4], mese)
    totali_controllo = ingest.carica_totali_controllo()

    rapporto = valida.valida(mese, config, movimenti, totali_controllo)
    if rapporto["bloccato"]:
        for errore in rapporto["errori"]:
            print(f"ERRORE   {errore['controllo']}: {errore['esito']}")
        print("Report NON rigenerati: errori bloccanti di validazione.")
        return 1

    esito_ribalta = ribalta.ribalta(movimenti, config)
    percorsi = genera_report(mese, config, movimenti, esito_ribalta,
                             budget, totali_controllo, rapporto)
    print(f"Report rigenerati in {percorsi[0].parent}:")
    for percorso in percorsi:
        print(f"  - {percorso.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
